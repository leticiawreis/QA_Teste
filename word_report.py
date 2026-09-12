"""Questionário e preenchimento do modelo de relatório QA em Word."""

import json
import re
import sys
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QTextEdit, QPushButton,
    QComboBox, QMessageBox, QFileDialog,
)


TEMPLATE = (Path(sys.executable).resolve().parent if getattr(sys, "frozen", False)
            else Path(__file__).resolve().parent) / "Template_QA_RPA_Geral.docx"
PLACEHOLDER = re.compile(r"\[[^\[\]]+\]")


@dataclass(frozen=True)
class Field:
    number: int
    section: str
    label: str


def _blocks(document):
    """Percorre parágrafos e tabelas na ordem em que aparecem no Word."""
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def _paragraphs(document):
    from docx.text.paragraph import Paragraph
    for block in _blocks(document):
        if isinstance(block, Paragraph):
            yield block, None
        else:
            for row in block.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        yield paragraph, block


def fields_from_template(template=TEMPLATE):
    document = Document(template)
    fields = []
    section = "Identificação do relatório"
    for paragraph, table in _paragraphs(document):
        if table is not None and len(table.rows) == 1 and len(table.columns) == 2:
            left = table.cell(0, 0).text.strip()
            right = table.cell(0, 1).text.strip()
            if left.isdigit() and right:
                section = f"{left}. {right}"
        for match in PLACEHOLDER.finditer(paragraph.text):
            label = match.group()[1:-1].strip()
            fields.append(Field(len(fields), section, label))
    return fields


def _replace_in_paragraph(paragraph, values):
    """Troca placeholders preservando o estilo do primeiro run de cada um."""
    for value in values:
        match = PLACEHOLDER.search(paragraph.text)
        if not match:
            return
        runs = paragraph.runs
        positions = []
        for run_index, run in enumerate(runs):
            positions.extend((run_index, offset) for offset in range(len(run.text)))
        if not positions:
            paragraph.text = paragraph.text[:match.start()] + value + paragraph.text[match.end():]
            continue
        start_run, start_offset = positions[match.start()]
        end_run, end_offset = positions[match.end() - 1]
        if start_run == end_run:
            run = runs[start_run]
            run.text = run.text[:start_offset] + value + run.text[end_offset + 1:]
        else:
            first = runs[start_run]
            last = runs[end_run]
            first.text = first.text[:start_offset] + value
            for run in runs[start_run + 1:end_run]:
                run.text = ""
            last.text = last.text[end_offset + 1:]


def generate_report(answers, output, template=TEMPLATE):
    document = Document(template)
    number = 0
    for paragraph, _ in _paragraphs(document):
        count = len(PLACEHOLDER.findall(paragraph.text))
        values = [answers.get(str(number + i), "").strip() or "Não informado"
                  for i in range(count)]
        _replace_in_paragraph(paragraph, values)
        number += count
    # Textos de orientação do modelo não fazem parte do relatório final.
    guidance = (
        "Use este documento como base", "Duplique o bloco abaixo",
        "Use esta seção para", "Use esta seção quando", "Registre itens como",
        "Use para divergências", "Registre somente o que foi realmente executado",
        "Adapte ao robô.", "Dica: evite colocar senhas",
    )
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(guidance):
            paragraph._element.getparent().remove(paragraph._element)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return output


class ReportWizard(QWidget):
    """Uma pergunta por vez, com rascunho para continuar depois."""

    def __init__(self, qa_dir):
        super().__init__()
        self.qa_dir = Path(qa_dir)
        self.draft = self.qa_dir / "06 - Relatórios Mensais" / "rascunho_relatorio_qa.json"
        self.fields = fields_from_template()
        self.answers = self._load_draft()
        self.index = 0

        layout = QVBoxLayout(self)
        title = QLabel("📝 Gerar relatório QA em Word")
        title.setObjectName("title")
        layout.addWidget(title)

        self.section = QComboBox()
        self.sections = list(dict.fromkeys(field.section for field in self.fields))
        self.section.addItems(self.sections)
        self.section.currentIndexChanged.connect(self._jump_section)
        layout.addWidget(self.section)

        self.progress = QLabel()
        self.question = QLabel()
        self.question.setWordWrap(True)
        self.answer = QTextEdit()
        self.answer.setPlaceholderText("Digite a informação para este campo. Deixe vazio se ainda não souber.")
        layout.addWidget(self.progress)
        layout.addWidget(self.question)
        layout.addWidget(self.answer)

        buttons = QHBoxLayout()
        previous = QPushButton("← Anterior")
        previous.clicked.connect(lambda: self._move(-1))
        next_button = QPushButton("Próximo →")
        next_button.clicked.connect(lambda: self._move(1))
        save = QPushButton("Salvar rascunho")
        save.clicked.connect(self._save_draft)
        generate = QPushButton("Gerar Word")
        generate.clicked.connect(self._generate)
        for button in (previous, next_button, save, generate):
            buttons.addWidget(button)
        layout.addLayout(buttons)
        self._show_field()

    def _load_draft(self):
        if self.draft.exists():
            try:
                data = json.loads(self.draft.read_text(encoding="utf-8"))
                if isinstance(data, dict):
                    return {str(key): str(value) for key, value in data.items()}
            except (OSError, ValueError):
                pass
        return {}

    def _remember(self):
        self.answers[str(self.index)] = self.answer.toPlainText().strip()

    def _show_field(self):
        field = self.fields[self.index]
        self.section.blockSignals(True)
        self.section.setCurrentIndex(self.sections.index(field.section))
        self.section.blockSignals(False)
        self.progress.setText(f"Campo {self.index + 1} de {len(self.fields)} · {field.section}")
        self.question.setText(f"O que colocar em: {field.label}?")
        self.answer.setPlainText(self.answers.get(str(self.index), ""))
        self.answer.setFocus()

    def _move(self, delta):
        self._remember()
        self.index = max(0, min(len(self.fields) - 1, self.index + delta))
        self._show_field()

    def _jump_section(self, section_index):
        self._remember()
        section = self.sections[section_index]
        self.index = next(i for i, field in enumerate(self.fields) if field.section == section)
        self._show_field()

    def _save_draft(self):
        self._remember()
        self.draft.parent.mkdir(parents=True, exist_ok=True)
        self.draft.write_text(json.dumps(self.answers, ensure_ascii=False, indent=2), encoding="utf-8")
        QMessageBox.information(self, "Relatório QA", f"Rascunho salvo em:\n{self.draft}")

    def _generate(self):
        self._remember()
        if not self.answers.get("0", "").strip():
            QMessageBox.warning(self, "Relatório QA", "Preencha primeiro o nome do robô.")
            return
        self.draft.parent.mkdir(parents=True, exist_ok=True)
        self.draft.write_text(json.dumps(self.answers, ensure_ascii=False, indent=2), encoding="utf-8")
        suggested = self.qa_dir / "06 - Relatórios Mensais" / "Relatorio_QA_RPA.docx"
        selected, _ = QFileDialog.getSaveFileName(self, "Salvar relatório Word", str(suggested), "Word (*.docx)")
        if not selected:
            return
        try:
            result = generate_report(self.answers, selected)
        except (OSError, ValueError) as error:
            QMessageBox.critical(self, "Erro ao gerar Word", str(error))
            return
        QMessageBox.information(self, "Relatório QA", f"Word criado em:\n{result}")
